"""A plain Excel copy of each backup, for the owner to open without the app.

It is built from the backup's own database file, so it matches the .sqlite3
file with the same name exactly, and it never locks the live database while
bills are being saved. Kept simple: one row per bill (with the items written
out in one cell), stock with the quantity sold, bookings and expenses.
Sheets: About, Bills, Stock, Bookings, Expenses.
"""

import sqlite3
from collections import defaultdict
from datetime import date, datetime, timezone as dt_timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from django.conf import settings
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from cash.models import Expense
from sales.models import Booking, Invoice, PaymentMode

MONEY = "#,##0.00"
DATE = "DD-MM-YYYY"
TIME = "HH:MM"
BOLD = Font(bold=True)

DOC_TYPES = dict(Invoice.DocType.choices)
BILL_STATUS = dict(Invoice.Status.choices)
BOOKING_STATUS = dict(Booking.Status.choices)
MODES = dict(PaymentMode.choices)
EXPENSE_FOR = dict(Expense.Category.choices)


def _local(value) -> datetime | None:
    """Stored times are UTC; the sheet shows shop time."""
    if not value:
        return None
    moment = datetime.fromisoformat(str(value))
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=dt_timezone.utc)
    return moment.astimezone(ZoneInfo(settings.TIME_ZONE)).replace(tzinfo=None)


def _day(value) -> date | None:
    return date.fromisoformat(str(value)[:10]) if value else None


def _rupees(paise) -> str:
    return f"{paise / 100:,.2f}"


class _Sheet:
    """A streamed sheet: bold header row, frozen, money and dates formatted."""

    def __init__(self, book: Workbook, title: str, columns: list[tuple[str, int, str]]):
        self.ws = book.create_sheet(title)
        for index, (_, width, _) in enumerate(columns, start=1):
            self.ws.column_dimensions[get_column_letter(index)].width = width
        self.ws.freeze_panes = "A2"
        self.kinds = [kind for _, _, kind in columns]
        header = []
        for name, _, _ in columns:
            cell = WriteOnlyCell(self.ws, value=name)
            cell.font = BOLD
            header.append(cell)
        self.ws.append(header)

    def add(self, *values) -> None:
        row = []
        for value, kind in zip(values, self.kinds):
            if value is None or value == "":
                row.append(None)
            elif kind == "money":
                cell = WriteOnlyCell(self.ws, value=value / 100)
                cell.number_format = MONEY
                row.append(cell)
            elif kind in ("date", "time"):
                cell = WriteOnlyCell(self.ws, value=value)
                cell.number_format = DATE if kind == "date" else TIME
                row.append(cell)
            elif isinstance(value, str):
                text = ILLEGAL_CHARACTERS_RE.sub("", value)
                if text.startswith("="):
                    # Typed text such as "=..." must stay text, never become a formula.
                    cell = WriteOnlyCell(self.ws, value=text)
                    cell.data_type = "s"
                    row.append(cell)
                else:
                    row.append(text)
            else:
                row.append(value)
        self.ws.append(row)


def write_backup_excel(db_file: Path, target: Path, taken_at: datetime, title: str = "Backup") -> None:
    conn = sqlite3.connect(f"{Path(db_file).resolve().as_uri()}?mode=ro", uri=True)
    try:
        book = Workbook(write_only=True)
        _about(book, conn, taken_at, title)
        _bills(book, conn)
        _stock(book, conn)
        _bookings(book, conn)
        _expenses(book, conn)
        book.save(target)
    finally:
        conn.close()


def _about(book, conn, taken_at, title) -> None:
    shop = conn.execute("SELECT shop_name FROM core_shopsettings LIMIT 1").fetchone()
    bills, sales = conn.execute(
        "SELECT COUNT(*), COALESCE(SUM(total_paise), 0) FROM sales_invoice WHERE status = 'active'"
    ).fetchone()
    cancelled = conn.execute("SELECT COUNT(*) FROM sales_invoice WHERE status = 'cancelled'").fetchone()[0]
    products = conn.execute("SELECT COUNT(*) FROM catalog_product").fetchone()[0]
    sheet = _Sheet(book, "About", [("Backup", 22, "text"), ("", 70, "text")])
    sheet.add("Shop", shop[0] if shop else "Not set up yet")
    sheet.add("Type", title)
    sheet.add("Backup taken", taken_at.strftime("%d-%m-%Y %I:%M %p"))
    sheet.add("Bills", f"{bills} (sales ₹{_rupees(sales)}, cancelled bills not counted)")
    sheet.add("Cancelled bills", str(cancelled))
    sheet.add("Products", str(products))
    sheet.add("Restore", "To restore the app, use the .sqlite3 file with the same name. This Excel file is for reading.")


def _bills(book, conn) -> None:
    paid = defaultdict(list)
    for invoice_id, mode, amount in conn.execute(
        "SELECT invoice_id, mode, SUM(amount_paise) FROM sales_payment "
        "WHERE invoice_id IS NOT NULL AND kind = 'sale' GROUP BY invoice_id, mode ORDER BY invoice_id, mode"
    ):
        paid[invoice_id].append(f"{MODES.get(mode, mode)} {_rupees(amount)}")
    items = defaultdict(list)
    for invoice_id, name, qty, sale_unit, unit, factor in conn.execute(
        "SELECT invoice_id, product_name, qty, sale_unit, unit_label, units_per_sale_unit "
        "FROM sales_invoiceline ORDER BY invoice_id, line_no"
    ):
        unit_text = f"Bundle of {factor}" if sale_unit == "case" else unit
        items[invoice_id].append(f"{name} x {qty} {unit_text}")
    sheet = _Sheet(
        book,
        "Bills",
        [
            ("Bill no", 16, "text"), ("Date", 12, "date"), ("Time", 8, "time"), ("Type", 12, "text"),
            ("Status", 11, "text"), ("Customer", 22, "text"), ("Phone", 13, "text"), ("Items", 60, "text"),
            ("Items total", 13, "money"), ("Discount", 11, "money"), ("Adjustment", 12, "money"), ("GST", 10, "money"),
            ("Round off", 10, "money"), ("Total", 13, "money"), ("Advance used", 13, "money"), ("Paid by", 30, "text"),
            ("Balance due", 12, "money"), ("Booking", 14, "text"), ("Counter", 12, "text"), ("Billed by", 14, "text"),
        ],
    )
    rows = conn.execute(
        "SELECT i.id, i.number, i.business_date, i.created_at, i.doc_type, i.status, i.customer_name, i.customer_phone, "
        "i.subtotal_paise, i.discount_paise, i.adjustment_paise, i.cgst_paise + i.sgst_paise, i.round_off_paise, "
        "i.total_paise, i.advance_paise, i.balance_paise, b.number, i.counter_name, i.user_name "
        "FROM sales_invoice i LEFT JOIN sales_booking b ON b.id = i.booking_id ORDER BY i.created_at, i.id"
    )
    for (bill_id, number, day, created, doc_type, status, name, phone, subtotal, discount, adjustment, gst, round_off,
         total, advance, balance, booking, counter, user) in rows:
        at = _local(created)
        sheet.add(
            number, _day(day), at.replace(microsecond=0).time() if at else None, DOC_TYPES.get(doc_type, doc_type),
            BILL_STATUS.get(status, status),
            name, phone, "; ".join(items.get(bill_id, [])), subtotal, discount or None, adjustment or None, gst or None,
            round_off or None, total, advance or None, ", ".join(paid.get(bill_id, [])), balance or None, booking,
            counter, user,
        )


def _stock(book, conn) -> None:
    places = conn.execute(
        "SELECT id, name FROM inventory_location WHERE is_active = 1 OR id IN "
        "(SELECT location_id FROM inventory_stockbalance WHERE quantity <> 0) ORDER BY is_sale_location DESC, id"
    ).fetchall()
    held = defaultdict(dict)
    for product_id, location_id, quantity in conn.execute(
        "SELECT product_id, location_id, quantity FROM inventory_stockbalance"
    ):
        held[product_id][location_id] = quantity
    sold = dict(
        conn.execute(
            "SELECT l.product_id, SUM(l.base_qty) FROM sales_invoiceline l "
            "JOIN sales_invoice i ON i.id = l.invoice_id WHERE i.status = 'active' GROUP BY l.product_id"
        ).fetchall()
    )
    columns = [("Code", 9, "text"), ("Item", 30, "text"), ("Category", 16, "text"), ("Unit", 8, "text"), ("Price", 11, "money")]
    columns += [(name, 13, "number") for _, name in places]
    columns += [("Total stock", 12, "number"), ("Sold", 9, "number"), ("On sale", 9, "text")]
    sheet = _Sheet(book, "Stock", columns)
    rows = conn.execute(
        "SELECT p.id, p.code, p.name, c.name, p.unit_name, p.price_paise, p.is_active "
        "FROM catalog_product p JOIN catalog_category c ON c.id = p.category_id ORDER BY c.sort_order, c.name, p.code"
    )
    for product_id, code, name, category, unit, price, active in rows:
        counts = [held[product_id].get(place_id, 0) for place_id, _ in places]
        sheet.add(
            code, name, category, unit, price, *counts, sum(held[product_id].values()), sold.get(product_id, 0),
            "Yes" if active else "No",
        )


def _bookings(book, conn) -> None:
    advance = dict(
        conn.execute(
            "SELECT booking_id, SUM(amount_paise) FROM sales_payment "
            "WHERE booking_id IS NOT NULL AND kind = 'advance' GROUP BY booking_id"
        ).fetchall()
    )
    items = defaultdict(list)
    for booking_id, name, qty, sale_unit, unit, per_case in conn.execute(
        "SELECT l.booking_id, p.name, l.qty, l.sale_unit, p.unit_name, p.units_per_case "
        "FROM sales_bookingline l JOIN catalog_product p ON p.id = l.product_id ORDER BY l.booking_id, l.id"
    ):
        unit_text = f"bundle of {per_case}" if sale_unit == "case" else unit
        items[booking_id].append(f"{name} x {qty} {unit_text}")
    sheet = _Sheet(
        book,
        "Bookings",
        [
            ("Booking no", 14, "text"), ("Booked on", 12, "date"), ("Deliver on", 12, "date"), ("Status", 11, "text"),
            ("Customer", 22, "text"), ("Phone", 13, "text"), ("Items", 50, "text"), ("Estimate", 12, "money"),
            ("Advance paid", 13, "money"), ("Address", 30, "text"), ("Note", 24, "text"),
        ],
    )
    rows = conn.execute(
        "SELECT id, number, created_at, delivery_date, status, customer_name, customer_phone, estimate_paise, address, note "
        "FROM sales_booking ORDER BY created_at, id"
    )
    for booking_id, number, created, deliver, status, name, phone, estimate, address, note in rows:
        at = _local(created)
        sheet.add(
            number, at.date() if at else None, _day(deliver), BOOKING_STATUS.get(status, status), name, phone,
            "; ".join(items.get(booking_id, [])), estimate or None, advance.get(booking_id), address, note,
        )


def _expenses(book, conn) -> None:
    sheet = _Sheet(
        book,
        "Expenses",
        [
            ("Date", 12, "date"), ("For", 14, "text"), ("Amount", 12, "money"), ("From cash drawer", 16, "text"),
            ("Note", 30, "text"), ("Entered by", 14, "text"),
        ],
    )
    rows = conn.execute(
        "SELECT e.business_date, e.category, e.amount_paise, e.paid_from_drawer, e.note, u.display_name "
        "FROM cash_expense e LEFT JOIN core_user u ON u.id = e.user_id ORDER BY e.business_date, e.created_at, e.id"
    )
    for day, category, amount, from_drawer, note, user in rows:
        sheet.add(_day(day), EXPENSE_FOR.get(category, category), amount, "Yes" if from_drawer else "No", note, user)
