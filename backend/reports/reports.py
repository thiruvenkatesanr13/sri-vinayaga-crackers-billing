"""Report tables for the Reports page and their Excel downloads.

Cancelled bills are never counted as sales. Each report returns columns +
rows + totals, so the screen and the Excel file always show the same numbers.
"""

import io
from collections import defaultdict
from datetime import date, timedelta

from django.db.models import Count, Sum
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from cash.models import DaySession, Expense
from catalog.models import Category
from core.errors import BusinessError
from core.models import ShopSettings
from sales.models import Invoice, InvoiceLine, Payment, PaymentMode

MODES = [(m.value, m.label) for m in PaymentMode]


def money(key, label):
    return {"key": key, "label": label, "type": "money"}


def num(key, label):
    return {"key": key, "label": label, "type": "int"}


def text(key, label):
    return {"key": key, "label": label, "type": "text"}


class Filters:
    def __init__(self, params, owner: bool):
        from django.utils import timezone
        from django.utils.dateparse import parse_date

        today = timezone.localdate()
        self.start = parse_date(params.get("from") or "") or today
        self.end = parse_date(params.get("to") or "") or self.start
        if self.end < self.start:
            raise BusinessError("The 'to' date is before the 'from' date.")
        if (self.end - self.start).days > 400:
            raise BusinessError("Choose a date range of up to 400 days.")
        self.category = None
        if params.get("category"):
            self.category = Category.objects.filter(pk=params["category"]).first()
        self.counter = params.get("counter") or None
        self.user_id = params.get("user") or None
        self.mode = params.get("mode") or None
        if self.mode and self.mode not in PaymentMode.values:
            raise BusinessError("Unknown payment mode.")
        self.owner = owner

    def invoices(self):
        rows = Invoice.objects.filter(
            status=Invoice.Status.ACTIVE, business_date__gte=self.start, business_date__lte=self.end
        )
        if self.counter:
            rows = rows.filter(counter_name=self.counter)
        if self.user_id:
            rows = rows.filter(user_id=self.user_id)
        if self.mode:
            rows = rows.filter(pk__in=Payment.objects.filter(mode=self.mode).values("invoice_id"))
        return rows

    def lines(self):
        rows = InvoiceLine.objects.filter(invoice__in=self.invoices())
        if self.category:
            rows = rows.filter(category_name=self.category.name)
        return rows

    def describe(self) -> str:
        parts = [f"{self.start:%d %b %Y}" + (f" to {self.end:%d %b %Y}" if self.end != self.start else "")]
        if self.category:
            parts.append(f"category {self.category.name}")
        if self.counter:
            parts.append(self.counter)
        if self.mode:
            parts.append(f"paid by {PaymentMode(self.mode).label}")
        return ", ".join(parts)


def _net_lines(filters: Filters):
    """Each line's share of the bill after discount and before GST."""
    for line in filters.lines().values(
        "product_code",
        "product_name",
        "category_name",
        "base_qty",
        "amount_paise",
        "unit_cost_paise",
        "invoice__taxable_paise",
        "invoice__subtotal_paise",
    ):
        subtotal = line["invoice__subtotal_paise"] or 0
        net = line["amount_paise"] * line["invoice__taxable_paise"] // subtotal if subtotal else 0
        yield line, net


def sales_summary(filters: Filters) -> dict:
    rows = []
    grouped = (
        filters.invoices()
        .values("business_date")
        .annotate(
            bills=Count("id", distinct=True),
            gross=Sum("subtotal_paise"),
            discount=Sum("discount_paise"),
            cgst=Sum("cgst_paise"),
            sgst=Sum("sgst_paise"),
            round_off=Sum("round_off_paise"),
            total=Sum("total_paise"),
        )
        .order_by("business_date")
    )
    for g in grouped:
        rows.append(
            {
                "date": g["business_date"].isoformat(),
                "bills": g["bills"],
                "gross_paise": g["gross"] or 0,
                "discount_paise": g["discount"] or 0,
                "gst_paise": (g["cgst"] or 0) + (g["sgst"] or 0),
                "round_off_paise": g["round_off"] or 0,
                "total_paise": g["total"] or 0,
            }
        )
    keys = ("bills", "gross_paise", "discount_paise", "gst_paise", "round_off_paise", "total_paise")
    totals = {k: sum(r[k] for r in rows) for k in keys}
    cancelled = Invoice.objects.filter(
        status=Invoice.Status.CANCELLED, business_date__gte=filters.start, business_date__lte=filters.end
    ).count()
    return {
        "title": "Sales summary",
        "columns": [
            text("date", "Date"),
            num("bills", "Bills"),
            money("gross_paise", "Items total"),
            money("discount_paise", "Discount"),
            money("gst_paise", "GST"),
            money("round_off_paise", "Round off"),
            money("total_paise", "Net sales"),
        ],
        "rows": rows,
        "totals": {"date": "Total", **totals},
        "notes": [f"{cancelled} cancelled bill(s) in this period are not counted."] if cancelled else [],
    }


def _grouped_sales(filters: Filters, key: str) -> dict:
    groups: dict[str, dict] = {}
    for line, net in _net_lines(filters):
        name = line["category_name"] if key == "category" else f"{line['product_code']}"
        entry = groups.setdefault(
            name,
            {
                "code": line["product_code"],
                "name": line["product_name"],
                "category": line["category_name"],
                "qty": 0,
                "amount_paise": 0,
                "net_paise": 0,
            },
        )
        entry["qty"] += line["base_qty"]
        entry["amount_paise"] += line["amount_paise"]
        entry["net_paise"] += net
    return groups


def product_sales(filters: Filters) -> dict:
    rows = sorted(_grouped_sales(filters, "product").values(), key=lambda r: -r["amount_paise"])
    return {
        "title": "Product-wise sales",
        "columns": [
            text("code", "Code"),
            text("name", "Product"),
            text("category", "Category"),
            num("qty", "Qty sold"),
            money("amount_paise", "Items total"),
            money("net_paise", "After discount"),
        ],
        "rows": rows,
        "totals": {
            "code": "Total",
            "qty": sum(r["qty"] for r in rows),
            "amount_paise": sum(r["amount_paise"] for r in rows),
            "net_paise": sum(r["net_paise"] for r in rows),
        },
        "notes": ["Qty is in each product's unit (a bundle counts as the pieces in it). GST is not included."],
    }


def category_sales(filters: Filters) -> dict:
    groups = _grouped_sales(filters, "category")
    rows = sorted(
        (
            {"category": k, "qty": v["qty"], "amount_paise": v["amount_paise"], "net_paise": v["net_paise"]}
            for k, v in groups.items()
        ),
        key=lambda r: -r["amount_paise"],
    )
    total_net = sum(r["net_paise"] for r in rows)
    for r in rows:
        r["share_bp"] = r["net_paise"] * 10000 // total_net if total_net else 0
    return {
        "title": "Category-wise sales",
        "columns": [
            text("category", "Category"),
            num("qty", "Qty sold"),
            money("amount_paise", "Items total"),
            money("net_paise", "After discount"),
            {"key": "share_bp", "label": "Share", "type": "percent"},
        ],
        "rows": rows,
        "totals": {
            "category": "Total",
            "qty": sum(r["qty"] for r in rows),
            "amount_paise": sum(r["amount_paise"] for r in rows),
            "net_paise": total_net,
        },
        "notes": [],
    }


def profit(filters: Filters) -> dict:
    if not filters.owner:
        from rest_framework.exceptions import PermissionDenied

        raise PermissionDenied("Only the owner can see profit.")
    groups: dict[str, dict] = {}
    missing: dict[str, dict] = {}
    for line, net in _net_lines(filters):
        target = missing if line["unit_cost_paise"] is None else groups
        entry = target.setdefault(
            line["product_code"],
            {"code": line["product_code"], "name": line["product_name"], "qty": 0, "net_paise": 0, "cost_paise": 0},
        )
        entry["qty"] += line["base_qty"]
        entry["net_paise"] += net
        if line["unit_cost_paise"] is not None:
            entry["cost_paise"] += line["unit_cost_paise"] * line["base_qty"]
    rows = []
    for entry in groups.values():
        entry["profit_paise"] = entry["net_paise"] - entry["cost_paise"]
        entry["margin_bp"] = entry["profit_paise"] * 10000 // entry["net_paise"] if entry["net_paise"] else 0
        rows.append(entry)
    rows.sort(key=lambda r: -r["profit_paise"])
    net_total = sum(r["net_paise"] for r in rows)
    profit_total = sum(r["profit_paise"] for r in rows)
    notes = []
    if missing:
        missing_sales = sum(m["net_paise"] for m in missing.values())
        names = ", ".join(sorted(m["name"] for m in missing.values())[:12])
        notes.append(
            f"{len(missing)} product(s) have no purchase cost, so their sales "
            f"(Rs {missing_sales / 100:,.2f}) are left out of profit: {names}"
            + (" ..." if len(missing) > 12 else "")
            + ". Add their cost on the Products page or enter a purchase."
        )
    return {
        "title": "Profit",
        "columns": [
            text("code", "Code"),
            text("name", "Product"),
            num("qty", "Qty sold"),
            money("net_paise", "Sales (after discount, before GST)"),
            money("cost_paise", "Purchase cost"),
            money("profit_paise", "Profit"),
            {"key": "margin_bp", "label": "Margin", "type": "percent"},
        ],
        "rows": rows,
        "totals": {
            "code": "Total",
            "qty": sum(r["qty"] for r in rows),
            "net_paise": net_total,
            "cost_paise": sum(r["cost_paise"] for r in rows),
            "profit_paise": profit_total,
            "margin_bp": profit_total * 10000 // net_total if net_total else 0,
        },
        "notes": notes,
        "missing_cost": sorted(missing.values(), key=lambda m: -m["net_paise"]),
    }


def daily_summary(filters: Filters) -> dict:
    days = []
    current = filters.start
    while current <= filters.end:
        days.append(current)
        current += timedelta(days=1)
    sales = {
        r["business_date"]: r
        for r in Invoice.objects.filter(
            status=Invoice.Status.ACTIVE, business_date__gte=filters.start, business_date__lte=filters.end
        )
        .values("business_date")
        .annotate(bills=Count("id"), total=Sum("total_paise"))
    }
    payments: dict[date, dict] = defaultdict(lambda: defaultdict(int))
    for r in (
        Payment.objects.filter(business_date__gte=filters.start, business_date__lte=filters.end)
        .values("business_date", "mode")
        .annotate(total=Sum("amount_paise"))
    ):
        payments[r["business_date"]][r["mode"]] = r["total"] or 0
    expenses = {
        r["business_date"]: r["total"] or 0
        for r in Expense.objects.filter(business_date__gte=filters.start, business_date__lte=filters.end)
        .values("business_date")
        .annotate(total=Sum("amount_paise"))
    }
    sessions = {
        s.business_date: s
        for s in DaySession.objects.filter(business_date__gte=filters.start, business_date__lte=filters.end)
    }
    rows = []
    for day in days:
        session = sessions.get(day)
        row = {
            "date": day.isoformat(),
            "bills": sales.get(day, {}).get("bills", 0),
            "total_paise": sales.get(day, {}).get("total", 0) or 0,
            "expenses_paise": expenses.get(day, 0),
            "opening_paise": session.opening_cash_paise if session else None,
            "expected_paise": session.expected_cash_paise if session else None,
            "counted_paise": session.counted_cash_paise if session else None,
            "difference_paise": session.difference_paise if session else None,
        }
        for mode, _label in MODES:
            row[f"{mode}_paise"] = payments[day].get(mode, 0)
        rows.append(row)
    mode_columns = [money(f"{m}_paise", label) for m, label in MODES if not filters.mode or m == filters.mode]
    sum_keys = ["bills", "total_paise", "expenses_paise"] + [c["key"] for c in mode_columns]
    return {
        "title": "Daily summary",
        "columns": [
            text("date", "Date"),
            num("bills", "Bills"),
            money("total_paise", "Net sales"),
            *mode_columns,
            money("expenses_paise", "Expenses"),
            money("opening_paise", "Opening cash"),
            money("expected_paise", "Expected cash"),
            money("counted_paise", "Counted cash"),
            money("difference_paise", "Cash difference"),
        ],
        "rows": rows,
        "totals": {"date": "Total", **{k: sum(r[k] for r in rows) for k in sum_keys}},
        "notes": ["Payment columns include booking advances and money returned for cancelled bills."],
    }


REPORTS = {
    "sales-summary": sales_summary,
    "products": product_sales,
    "categories": category_sales,
    "profit": profit,
    "daily": daily_summary,
}


def build(kind: str, params, owner: bool) -> dict:
    if kind not in REPORTS:
        raise BusinessError("Unknown report.")
    filters = Filters(params, owner)
    report = REPORTS[kind](filters)
    report["kind"] = kind
    report["period"] = filters.describe()
    return report


def to_xlsx(report: dict) -> bytes:
    shop = ShopSettings.get()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report["title"][:31]

    def put(row, col, value, number_format=None, bold=False):
        cell = sheet.cell(row=row, column=col, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        if number_format:
            cell.number_format = number_format
        if bold:
            cell.font = Font(bold=True)
        return cell

    put(1, 1, shop.shop_name, bold=True)
    put(2, 1, f"{report['title']}: {report['period']}")
    columns = report["columns"]
    for col, column in enumerate(columns, start=1):
        put(4, col, column["label"], bold=True).alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[sheet.cell(row=4, column=col).column_letter].width = max(
            12, min(40, len(column["label"]) + 4)
        )

    def write_row(row_no, row, bold=False):
        for col, column in enumerate(columns, start=1):
            value = row.get(column["key"])
            if value is None:
                continue
            if column["type"] == "money":
                put(row_no, col, value / 100, "#,##0.00", bold)
            elif column["type"] == "percent":
                put(row_no, col, value / 10000, "0.0%", bold)
            else:
                put(row_no, col, value, None, bold)

    row_no = 5
    for row in report["rows"]:
        write_row(row_no, row)
        row_no += 1
    if report.get("totals"):
        write_row(row_no, report["totals"], bold=True)
        row_no += 1
    for note in report.get("notes", []):
        row_no += 1
        put(row_no, 1, note)
    sheet.freeze_panes = "A5"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
